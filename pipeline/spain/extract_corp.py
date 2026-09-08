"""Who generates corporate tax: AEAT breakdown by company type (table 8.5).

The latest one or two columns are headed "2024 (p)" / "2025(p)" — provisional.
They are read like any other year and the flag is carried, because a header that
is not a bare integer was being dropped silently and the console then showed the
previous year's table under the new year's heading."""
import openpyxl, json, re
wb=openpyxl.load_workbook('who/Cuadros_IART25_es_es.xlsx',data_only=True)
ws=wb['8.5']
years={}   # year -> (column, provisional)
for c in range(3,70):
    v=str(ws.cell(6,c).value or '').strip()
    m=re.match(r'^((?:19|20)\d\d)\s*(\(p\))?$', v)
    if m and int(m.group(1))>1990: years[m.group(1)]=(c, bool(m.group(2)))
BLOCKS={'total':9,'groups':36,'standalone':63}
FIELDS={'profit':'Resultado contable positivo','base':'Base imponible positiva',
        'tax':'Cuota líquida positiva','rateBase':'Tipo efectivo sobre BI (%)',
        'rateProfit':'Tipo efectivo sobre RC>0 (%)','exempt':'Exención por doble imposición (1)',
        'losses':'Compensación de bases negativas de períodos anteriores'}
def findrow(start, label, span=30):
    for r in range(start, start+span):
        b=ws.cell(r,2).value
        if b and str(b).strip().startswith(label[:34]): return r
    return None
out={}
for y,(col,prov) in years.items():
    rec={'prov':prov}
    for bk,start in BLOCKS.items():
        d={}
        for f,lab in FIELDS.items():
            r=findrow(start,lab)
            v=ws.cell(r,col).value if r else None
            d[f]= (round(float(v),2) if isinstance(v,(int,float)) else None)
        rec[bk]=d
    # a column with no figures (2025 today) is not a year
    if rec['total']['profit'] is None: continue
    out[y]=rec
json.dump(out,open('corp_types.json','w'))
y=sorted(k for k,v in out.items() if v['total']['profit'] is not None)[-1]; r=out[y]
print(f"SPAIN {y} — corporate tax by company type (€bn, AEAT table 8.5)\n")
print("  type          profit    tax base      tax   eff.on base  eff.on profit")
for k,lab in [('total','ALL companies'),('groups','Consolidated groups'),('standalone','Standalone cos')]:
    d=r[k]
    print(f"  {lab:<20}{d['profit']/1000:>7.1f}{d['base']/1000:>11.1f}{d['tax']/1000:>10.1f}"
          f"{str(d['rateBase'])+'%':>12}{str(d['rateProfit'])+'%':>14}")
g,s,t=r['groups'],r['standalone'],r['total']
print(f"\n  groups make {g['profit']/t['profit']*100:.0f}% of profits but pay {g['tax']/t['tax']*100:.0f}% of the tax")
print(f"  standalone make {s['profit']/t['profit']*100:.0f}% of profits and pay {s['tax']/t['tax']*100:.0f}% of the tax")
print(f"  double-taxation exemption removes €{abs(g['exempt'])/1000:.1f}bn from group profits before tax")
