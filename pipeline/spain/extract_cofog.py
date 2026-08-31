import openpyxl, json, re, glob, unicodedata
def norm(s):
    s=unicodedata.normalize('NFD',str(s)).encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z]','',s)
# sheet title (A1) -> our region id
NAME2ID={'andalucia':'01','aragon':'02','asturias':'03','baleares':'04','canarias':'05',
'cantabria':'06','castillayleon':'07','castillalamancha':'08','cataluna':'09',
'comunidadvalenciana':'10','comunitatvalenciana':'10','valencia':'10','extremadura':'11','galicia':'12','madrid':'13',
'regiondemurcia':'14','murcia':'14','comunidadforaldenavarra':'15','navarra':'15',
'paisvasco':'16','larioja':'17'}
ECON=['D.1','P.2','D.3','D.4','D.62','D.632','D.7','P.5','D.9']
ECON_LABEL={'D.1':'Salaries','P.2':'Goods & services','D.3':'Subsidies','D.4':'Property income (interest)',
'D.62':'Social benefits (cash)','D.632':'Social transfers in kind','D.7':'Other current transfers',
'P.5':'Investment','D.9':'Capital transfers'}
out={}   # year -> region -> {div: total, econ:{}}
agg={}   # year -> {div: total} for "Administración regional" (all CCAA)
for f in sorted(glob.glob('cofog_ccaa/d_*.xlsx')):
    year=re.search(r'd_(\d{4})',f).group(1)
    wb=openpyxl.load_workbook(f,data_only=True)
    out[year]={}
    for sn in wb.sheetnames:
        if sn=='Indice': continue
        ws=wb[sn]
        title=ws.cell(1,1).value or ''
        key=norm(title)
        # header row: find the row whose col A equals the year label e.g. "2024(P)" or "2024"
        hdr=None
        for r in range(5,12):
            v=ws.cell(r,1).value
            if v and re.match(r'^\s*'+year, str(v)): hdr=r;break
        if hdr is None: continue
        # map columns -> division code from header text
        cols={}
        for c in range(3,30):
            h=ws.cell(hdr,c).value
            if not h: continue
            h=str(h).strip()
            m=re.search(r'(\d{2})\.',h) or re.match(r'^\s*(\d{2})',h)
            if m: cols[c]=m.group(1)
            elif 'TOTAL' in h.upper(): cols[c]='TOT'
        if not cols: continue
        # data rows
        totrow=None; econ={}
        for r in range(hdr+1,hdr+40):
            a=str(ws.cell(r,1).value or '').strip()
            b=str(ws.cell(r,2).value or '').strip().upper()
            if 'GASTO TOTAL' in b: totrow=r
            if a in ECON: econ[a]=r
        if totrow is None: continue
        def rowvals(r):
            d={}
            for c,code in cols.items():
                v=ws.cell(r,c).value
                if isinstance(v,(int,float)): d[code]=round(float(v))
            return d
        rec={'div':rowvals(totrow),
             'econ':{k:rowvals(v).get('TOT',0) for k,v in econ.items()}}
        if key in NAME2ID:
            out[year][NAME2ID[key]]=rec
        elif 'administracionregional' in key or 'administracinregional' in key:
            agg[year]=rec
json.dump({'regions':out,'aggregate':agg,'econLabels':ECON_LABEL},open('cofog_regional.json','w'))
yrs=sorted(out)
print("years:",yrs[0],"-",yrs[-1])
for y in [yrs[0],'2020',yrs[-1]]:
    if y in out:
        n=len(out[y]); tot=sum(r['div'].get('TOT',0) for r in out[y].values())
        a=agg.get(y,{}).get('div',{}).get('TOT')
        print(f"  {y}: {n} regions, sum {tot/1000:.1f}bn"+(f", published aggregate {a/1000:.1f}bn, diff {(tot-a)/1000:+.1f}bn" if a else ""))
print("\n2024 top spenders (€bn):")
for rid,rec in sorted(out['2024'].items(), key=lambda kv:-kv[1]['div'].get('TOT',0))[:6]:
    d=rec['div']
    print(f"  {rid}  tot {d.get('TOT',0)/1000:6.1f}   health {d.get('07',0)/1000:5.1f}   educ {d.get('09',0)/1000:5.1f}")
