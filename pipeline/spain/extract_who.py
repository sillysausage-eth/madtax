"""Who generates income tax: AEAT decile distribution, national series."""
import openpyxl, json
wb=openpyxl.load_workbook('who/DistribucionesIRPF.xlsx',data_only=True)
ws=wb['datos']
COL=dict(year=1, scope=2, ccaa=3, limit=4, band=5, n=6, total=7,
         work=8, capMob=9, capInm=10, biz=11, gains=12, imputed=13,
         tax=14, rate=15)
out={}
for r in range(2, ws.max_row+1):
    if str(ws.cell(r,COL['scope']).value or '').strip()!='Nacional': continue
    y=str(ws.cell(r,COL['year']).value or '').strip()
    band=str(ws.cell(r,COL['band']).value or '').strip()
    if not y or not band: continue
    g=lambda k: ws.cell(r,COL[k]).value
    num=lambda v: (round(float(v)) if isinstance(v,(int,float)) else None)
    lim=g('limit')
    out.setdefault(y,{})[band]={
        'limit': (round(float(lim)) if isinstance(lim,(int,float)) else None),
        'n': num(g('n')),
        'income': num(g('total')),
        'tax': num(g('tax')),
        'rate': (round(float(g('rate')),2) if isinstance(g('rate'),(int,float)) else None),
        'src': {k: num(g(k)) for k in ['work','capMob','capInm','biz','gains','imputed']}
    }
json.dump(out, open('irpf_deciles.json','w'))
yrs=sorted(out)
print(f"years {yrs[0]}-{yrs[-1]}, bands per year: {len(out[yrs[-1]])}")
y=yrs[-1]; d=out[y]
print(f"\nSPAIN {y} — income tax by decile (national)")
print("  band   upper €      taxpayers      income €bn    tax €bn   eff.rate   %tax")
tot=d.get('TOT')
order=[f'D{i:02d}' for i in range(1,11)]+['P99','P999','P9999']
for b in order:
    x=d.get(b)
    if not x: continue
    share = x['tax']/tot['tax']*100 if tot and tot['tax'] else 0
    lim = f"{x['limit']:>10,}" if x['limit'] else "         —"
    print(f"  {b:<6}{lim}  {x['n']:>12,}  {x['income']/1e9:>11.1f}  {x['tax']/1e9:>9.1f}  {str(x['rate'])+'%':>8}  {share:>5.1f}%")
print(f"  {'TOTAL':<6}{'':>10}  {tot['n']:>12,}  {tot['income']/1e9:>11.1f}  {tot['tax']/1e9:>9.1f}  {str(tot['rate'])+'%':>8}")
