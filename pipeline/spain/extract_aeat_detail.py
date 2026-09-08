"""Administrative detail AEAT publishes and national accounts do not:
   excise by product (table 5.1) and VAT by rate (table 8.7).

   Both are AEAT accrued state figures ("devengado"), so they do NOT equal the
   Eurostat ESA bucket for general government — the foral territories collect their
   own, and the ESA bucket is wider. We keep AEAT's own total so the gap is visible
   rather than papered over with a rescaling."""
import openpyxl, re, json, unicodedata

# The annexes of the Informe Anual de Recaudación Tributaria. AEAT renamed the
# file between editions (Cuadros_IART24.xlsx, then Cuadros_IART25_es_es.xlsx),
# so the name is not derived from the year.
WB = 'who/Cuadros_IART25_es_es.xlsx'
wb = openpyxl.load_workbook(WB, data_only=True)

def squash(s):
    s = unicodedata.normalize('NFKD', str(s or ''))
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]', '', s.lower())

def row_of(ws, label, after=1, before=None):
    """The first row at or after `after` whose column-B label squashes to `label`.

    Rows are found by label, never by position: the 2025 edition inserted the
    e-cigarette liquid line into table 5.1 and pushed coal, plastic and
    electricity down one row each, and re-ordered the adjustments block of 8.7.
    A position-based read would have silently filed coal as e-cigarettes."""
    before = before or ws.max_row
    for r in range(after, before + 1):
        if squash(ws.cell(r, 2).value) == label:
            return r
    raise SystemExit(f'{ws.title}: row "{label}" not found between {after} and {before}')

def year_cols(ws):
    """Year headers run left to right and then give way to a percent-change block;
    stop at the first non-year once we have started. '(p)' marks provisional."""
    out = {}
    for c in range(3, ws.max_column + 1):
        v = str(ws.cell(6, c).value or '').strip()
        m = re.match(r'^((?:19|20)\d\d)(\s*\(p\))?$', v)
        if m:
            out[m.group(1)] = (c, bool(m.group(2)))
        elif out:
            break
    return out

def num(v):
    return round(float(v)) if isinstance(v, (int, float)) else None

def block(ws, cols, total_row, rows, since=2012):
    """Rows that partition total_row. Returns {year: {'total':n,'prov':bool,'rows':[[k,v]]}}"""
    out = {}
    for y, (c, prov) in cols.items():
        if int(y) < since:
            continue
        tot = num(ws.cell(total_row, c).value)
        if tot is None:
            continue
        parts = []
        for r, k in rows:
            v = num(ws.cell(r, c).value)
            if v:
                parts.append([k, v])
        if not parts:
            continue
        s = sum(v for _, v in parts)
        if abs(s - tot) > 2:
            print(f"  MISMATCH {y}: parts {s} vs total {tot} ({s-tot:+d})")
        out[y] = {'total': tot, 'prov': prov,
                  'rows': sorted(parts, key=lambda x: -x[1])}
    return out

# ---- 5.1 excise duties accrued, by product ----
ws = wb['5.1']
tot = row_of(ws, 'impuestoespecialdevengadod2')
EXCISE_LINES = [
    ('alcoholybebidasderivadas', 'alcohol'), ('cerveza', 'beer'),
    ('productosintermedios', 'intermediate'), ('hidrocarburos', 'fuel'),
    ('laboresdeltabaco', 'tobacco'), ('liquidocigarrilloselectronicos', 'ecig'),
    ('carbon', 'coal'), ('envasesdeplasticonoreutilizables', 'plastic'),
    ('electricidad', 'electricity')]
excise = block(ws, year_cols(ws), tot, [
    (row_of(ws, lab, tot, tot + 12), k) for lab, k in EXCISE_LINES])

# ---- 8.7 VAT accrued, by rate. The rate split is published for the general regime
#      only; the special regimes and the foral adjustment sit outside it. ----
ws = wb['8.7']
cols = year_cols(ws)
accrued = row_of(ws, 'ivadevengadoenelperiodoivadivagivareaj')
general = row_of(ws, 'ivaenelregimengeneralivagab', accrued)
special = row_of(ws, 'ivaenlosregimenesespecialesivarecd', general)
# The general-regime rate lines sit between the two subtotals. The 2% line is the
# temporary rate on basic foods (Oct-Dec 2024); the 2024 edition labelled the
# same row "Tipo 2,5" and the 2025 edition corrected it to "Tipo 2".
RATE_LINES = [('tipo0', 'r0'), ('tipo25', 'r2'), ('tipo2', 'r2'), ('tiposuperreducido', 'rsuper'),
              ('tipo5', 'r5'), ('tipo75', 'r75'), ('tiporeducido', 'rreduced'),
              ('tipogeneral', 'rgeneral')]
rate_rows, seen = [], set()
for r in range(general + 1, special):
    lab = squash(ws.cell(r, 2).value)
    for pat, k in RATE_LINES:
        if lab == pat and k not in seen:
            rate_rows.append((r, k)); seen.add(k)
vat = block(ws, cols, general, rate_rows)
# context rows so the page can say what the rate split leaves out
foral_row = row_of(ws, 'ivacorrespondientealosterritoriosforalese', special)
other_row = row_of(ws, 'otrosf', special)
for y, (c, _) in cols.items():
    if y in vat:
        vat[y]['accrued'] = num(ws.cell(accrued, c).value)   # all accrued VAT
        vat[y]['special'] = num(ws.cell(special, c).value)   # special regimes
        vat[y]['foral']   = num(ws.cell(foral_row, c).value) # foral territories
        vat[y]['adjOther']= num(ws.cell(other_row, c).value) # other adjustments

json.dump({'excise': excise, 'vat': vat}, open('aeat_detail.json', 'w'))

ey, vy = sorted(excise)[-1], sorted(vat)[-1]
print(f"excise {sorted(excise)[0]}-{ey} · vat {sorted(vat)[0]}-{vy}")
print(f"\nEXCISE {ey} (AEAT accrued, total €{excise[ey]['total']/1000:.1f}bn"
      f"{' provisional' if excise[ey]['prov'] else ''})")
for k, v in excise[ey]['rows']:
    print(f"  {k:<13}{v/1000:>7.2f}bn{v/excise[ey]['total']*100:>7.1f}%")
print(f"\nVAT BY RATE {vy} (general regime, €{vat[vy]['total']/1000:.1f}bn"
      f"{' provisional' if vat[vy]['prov'] else ''})")
for k, v in vat[vy]['rows']:
    print(f"  {k:<13}{v/1000:>7.2f}bn{v/vat[vy]['total']*100:>7.1f}%")
print(f"  all accrued VAT €{vat[vy]['accrued']/1000:.1f}bn = general regime"
      f" + special €{vat[vy]['special']/1000:.1f}bn"
      f" + foral €{vat[vy]['foral']/1000:.1f}bn"
      f" + other €{vat[vy]['adjOther']/1000:.1f}bn")
