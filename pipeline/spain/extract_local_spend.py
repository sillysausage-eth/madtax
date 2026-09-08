"""Local government EXPENDITURE by autonomous community, from CONPREL.

The spending map shows what is spent in each territory, whoever spends it. The
regional tier comes from IGAE's per-community COFOG files; this is the other tier
that has a territory — town councils, provincial and island councils, comarcas,
metropolitan areas and the two autonomous cities — read from the same CONPREL
definitive liquidations `extract_local.py` reads the revenue side of.

Values in the source are THOUSANDS of euros; emitted here as € millions.

WHAT IS READ
------------
Tabla 3  Desglose de los gastos consolidados — every chapter and article of the
         consolidated local expenditure of the community, "Obligaciones
         Reconocidas Netas" (obligations recognised in the year).
Tabla 2  Desglose de los ingresos consolidados — only articles 45 and 75, the
         current and capital transfers received FROM the community government,
         "Derechos Reconocidos Netos".
Tabla 4  Clasificación económica y por programas — the same obligations by the
         six programme areas of Orden EHA/3565/2008 (0 debt, 1 basic public
         services, 2 social protection and promotion, 3 preferential public goods,
         4 economic action, 9 general). Financial chapters are removed column by
         column so the areas partition the same figure as chapters 1-7.

WHAT IS EXCLUDED, AND WHY
-------------------------
* Chapters 8 and 9 (financial assets, debt repayment). National accounts do not
  count them as expenditure and the national figure the map reconciles to is a
  national-accounts figure.
* Transfers to another tier of government — articles 42/72 (State), 43/73 (Social
  Security), 45/75 (community government), 46/76 (other local entities). That
  money is spent by the recipient and is already on the map or in the coin; in
  the Basque Country the Diputaciones Forales hand ~€12bn a year of concierto tax
  to the Basque Government, and counting it here as well would count it twice.
* Transfers received from the community government (income articles 45/75). The
  community books them as its own spending in the IGAE figure the map already
  carries, and the council books them again when it spends them. One leg comes
  out so the territory reads once.

Nothing is estimated. A community whose CONPREL table is published but zero on
every line (Navarre 2013, 2014) is returned as absent, exactly as the revenue
extractor does, and the merge names the gap.
"""
import glob, json, os, re
import openpyxl, xlrd
from regions import resolve, assert_complete

CITY_IDS = {'ceuta': '18', 'melilla': '19'}
TO_GOV = ('42', '43', '45', '46', '72', '73', '75', '76')
FROM_CA = ('45', '75')
NONFIN_CH = tuple(str(i) for i in range(1, 8))
FIN_CH = ('8', '9')
AREA_RE = re.compile(r'^[0-9]$')
AREA_LABELS = {}


def _norm(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip().lower()


class Sheet:
    """One table, addressed (row, col) zero-based, for either file format."""

    def __init__(self, wb, wanted):
        self.ws = None
        if isinstance(wb, openpyxl.Workbook):
            name = next((n for n in wb.sheetnames if n.strip() in wanted), None)
            if name is None:
                raise KeyError(wanted[0])
            ws = wb[name]
            self.nrows, self.ncols = ws.max_row, ws.max_column
            self.get = lambda r, c: ws.cell(r + 1, c + 1).value if r < self.nrows and c < self.ncols else None
        else:
            name = next((n for n in wb.sheet_names() if n.strip() in wanted), None)
            if name is None:
                raise KeyError(wanted[0])
            sh = wb.sheet_by_name(name)
            self.nrows, self.ncols = sh.nrows, sh.ncols
            self.get = lambda r, c: sh.cell_value(r, c) if r < sh.nrows and c < sh.ncols else None

    def region(self):
        for r in range(8):
            for c in range(4):
                v = self.get(r, c)
                rid = resolve(v)
                if rid:
                    return rid
                for k, cid in CITY_IDS.items():
                    if v and k in str(v).lower():
                        return cid
        return None

    def header_col(self, *needles):
        """The column whose header cell (rows 4-10, possibly split over two
        rows) contains every needle. Read by text, never by position: the two
        CONPREL formats are one column apart."""
        for c in range(self.ncols):
            text = _norm(' '.join(str(self.get(r, c) or '') for r in range(4, 11)))
            if all(n in text for n in needles):
                return c
        raise SystemExit(f'column {needles} not found')

    def code_col(self):
        for c in range(4):
            for r in range(4, 12):
                if _norm(self.get(r, c)).startswith('ctas'):
                    return c
        raise SystemExit('code column not found')

    def rows(self, col, label_col=None, labels=None, label_re=None):
        cc = self.code_col()
        out = {}
        for r in range(self.nrows):
            code = str(self.get(r, cc) or '').strip()
            if code.endswith('.0'):
                code = code[:-2]
            if not code:
                continue
            v = self.get(r, col)
            if isinstance(v, (int, float)):
                out.setdefault(code, float(v))
            if labels is not None and label_re and label_re.match(code):
                labels.setdefault(code, str(self.get(r, cc + 1) or '').strip())
        return out

    def total(self, col):
        cc = self.code_col()
        for r in range(self.nrows):
            if str(self.get(r, cc) or '').strip() == '' and 'total' in _norm(self.get(r, cc + 1)):
                v = self.get(r, col)
                if isinstance(v, (int, float)):
                    return float(v)
        raise SystemExit('total row not found')


def open_book(path):
    if path.endswith('.xlsx'):
        return openpyxl.load_workbook(path, data_only=True)
    return xlrd.open_workbook(path)


def read(path):
    wb = open_book(path)
    t3 = Sheet(wb, ('Tabla 3', 'Tabla 3.0'))
    rid = t3.region()
    if rid is None:
        raise SystemExit(f'could not resolve region from {path}')
    orn = t3.header_col('obligaciones', 'reconocidas')
    exp = t3.rows(orn)
    total = t3.total(orn)
    g = lambda c: exp.get(c, 0.0)

    # An all-zero published table is an absent figure, not a council that spent
    # nothing: Navarre 2013 and 2014.
    if total == 0 and all(g(c) == 0 for c in NONFIN_CH):
        return rid, None

    chapters = sum(g(c) for c in NONFIN_CH + FIN_CH)
    if abs(chapters - total) > 1:
        raise SystemExit(f'{path}: chapters {chapters} vs total {total}')
    nonfin = sum(g(c) for c in NONFIN_CH)
    fin = sum(g(c) for c in FIN_CH)
    to_gov = sum(g(c) for c in TO_GOV)

    t2 = Sheet(wb, ('Tabla 2', 'Tabla 2.0'))
    drn = t2.header_col('derechos', 'reconocidos')
    inc = t2.rows(drn)
    from_ca = sum(inc.get(c, 0.0) for c in FROM_CA)

    t4 = Sheet(wb, ('Tabla 4', 'Tabla 4.0'))
    tot_c = t4.header_col('total')
    act_c = t4.header_col('activos')
    pas_c = t4.header_col('pasivos')
    areas_tot = t4.rows(tot_c, labels=AREA_LABELS, label_re=AREA_RE)
    areas_act = t4.rows(act_c)
    areas_pas = t4.rows(pas_c)
    areas = {}
    for code, v in areas_tot.items():
        if AREA_RE.match(code):
            areas[code] = v - areas_act.get(code, 0.0) - areas_pas.get(code, 0.0)
    # The areas must partition the non-financial figure the chapters give, or
    # the file's layout has moved under us. Tabla 4 rounds independently of
    # Tabla 3, so the tolerance is €1k per area.
    if abs(sum(areas.values()) - nonfin) > len(areas):
        raise SystemExit(f'{path}: areas {sum(areas.values())} vs chapters 1-7 {nonfin}')

    m = lambda v: round(v / 1000)  # thousands -> millions
    return rid, {
        'total': m(total), 'nonfin': m(nonfin), 'fin': m(fin),
        'toGov': m(to_gov), 'fromCA': m(from_ca),
        'net': m(nonfin - to_gov - from_ca),
        'areas': {c: m(v) for c, v in sorted(areas.items())},
    }


out, empty = {}, {}
for f in sorted(glob.glob('local/ccaa/EL*C*.xls') + glob.glob('local/ccaa/EL*C*.xlsx')):
    year = re.search(r'EL(\d{4})C', os.path.basename(f)).group(1)
    rid, rec = read(f)
    out.setdefault(year, {})
    if rec:
        if rid in out[year]:
            raise SystemExit(f'{year}: region {rid} appears twice')
        out[year][rid] = rec
    else:
        empty.setdefault(year, []).append(rid)

for y, d in out.items():
    assert_complete(set(d) | set(empty.get(y, [])), f'CONPREL expenditure {y}')
    for cid in CITY_IDS.values():
        if cid not in d and cid not in empty.get(y, []):
            raise SystemExit(f'{y}: autonomous city {cid} missing')
for y, ids in sorted(empty.items()):
    print(f'NOTE {y}: CONPREL publishes an all-zero table for region(s) {ids} — carried as absent, not zero')

json.dump(out, open('local_spend.json', 'w'))
json.dump({'codes': sorted(AREA_LABELS), 'labels': AREA_LABELS},
          open('local_spend_areas.json', 'w'), ensure_ascii=False)

yrs = sorted(out)
print(f'years {yrs[0]}-{yrs[-1]}, territories per year: ' + ', '.join(f'{y}:{len(out[y])}' for y in yrs))
for c in sorted(AREA_LABELS):
    print(f'  area {c}: {AREA_LABELS[c]}')
y = yrs[-1]
S = lambda k: sum(r[k] for r in out[y].values())
print(f'\n{y} roll-up (€bn): obligations {S("total")/1000:.1f} · non-financial {S("nonfin")/1000:.1f} · '
      f'to other tiers {S("toGov")/1000:.1f} · from community govts {S("fromCA")/1000:.1f} · net on map {S("net")/1000:.1f}')
for rid, r in sorted(out[y].items(), key=lambda kv: -kv[1]['net'])[:5]:
    print(f'  {rid}  nonfin {r["nonfin"]/1000:6.2f}  toGov {r["toGov"]/1000:6.2f}  fromCA {r["fromCA"]/1000:5.2f}  net {r["net"]/1000:6.2f}')
